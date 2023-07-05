"""Python to Excel"""
import pandas as pd
from openpyxl import Workbook
from openpyxl.utils.dataframe import dataframe_to_rows
from openpyxl.chart import BarChart, Reference

pd.set_option("display.max_column",500)

FILENAME = "video_game_sales.csv"

raw_df = pd.read_csv(FILENAME)


def num_games(_df):
    """How many games per platform"""
    _games_df  = _df.groupby(["Platform"])["Name"].count().reset_index()
    return _games_df

games_df  = num_games(raw_df)

wb = Workbook()

ws1 = wb.active
ws1.column_dimensions['A'].width = 20
ws1.column_dimensions['B'].width = 20

chart1 = BarChart()
chart1.type = "col"
chart1.style = 10
chart1.title = "Video Games By Platform"
chart1.y_axis.title = '# of games'
chart1.x_axis.title = 'Platform'

for r in dataframe_to_rows(games_df, index=False, header=True):
    ws1.append(r)

last_row = len(ws1["A"])

data = Reference(ws1, min_col=2, min_row=1, max_col=2, max_row=last_row)
cats = Reference(ws1, min_col=1, min_row=2, max_col=1, max_row=last_row)

chart1.add_data(data, titles_from_data=True)
chart1.set_categories(cats)

chart1.shape = 4
ws1.add_chart(chart1, "F1")

wb.save("Week9PyToExcel.xlsx")

# What countries had the highest sales for N64?
# What were the TOP 5 games for the PC?